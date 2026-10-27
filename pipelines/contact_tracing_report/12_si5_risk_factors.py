"""
Step 12: supplementary table SI-5 -- risk factors for becoming a secondary case.
Reads tables/sar_risk_factors.csv and tables/sar_multivariable_or.csv (02_analysis.py). Writes:
  tables/si5_risk_factors.csv
  si5_risk_factors.docx        (formatted Word table, landscape, in the output folder)

For each factor level: k/n eligible pairs, SAR, crude risk ratio (vs reference level) and adjusted
odds ratio (single logistic regression including exposure type, setting, relationship, source
province, source outcome and month listed). Each estimate is shown with the 95% CI from bootstrap
resampling of source cases (used in the text) and the conventional CI (Wilson for SAR, Wald for
RR / OR) alongside.
"""
import os
import numpy as np
import pandas as pd
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

import json
from ctr_paths import T, D, OUT_ROOT, MONTH_ORDER, FU_WINDOW_END, fmt_day

rf = pd.read_csv(os.path.join(T, "sar_risk_factors.csv"))
mv = pd.read_csv(os.path.join(T, "sar_multivariable_or.csv")).set_index("term")
KEY = {"Exposure type": "contact_type_", "Setting": "case_circumstances_", "Relationship": "relationship_to_case_",
       "Province (source)": "src_province_", "Source outcome": "source_died_", "Month listed": "reg_month_"}
ORDER = ["Exposure type", "Setting", "Relationship", "Province (source)", "Month listed", "Source outcome", "Contacts listed by source"]
# MONTH_ORDER (May .. month of the cut) comes from ctr_paths
N_ELIG = json.load(open(os.path.join(D, "results.json")))["sar_eligible_pairs"]
SIZE_ORDER = ["1-5", "6-15", "16-50", ">50"]

ci = lambda a, b, d=1: f"{a:.{d}f}–{b:.{d}f}"
rows = []
for fac in ORDER:
    d = rf[rf.factor == fac].copy()
    if fac == "Month listed":
        d["o"] = d.level.map({m: i for i, m in enumerate(MONTH_ORDER)})
    elif fac == "Contacts listed by source":
        d["o"] = d.level.map({m: i for i, m in enumerate(SIZE_ORDER)})
    else:
        d["o"] = -d.n                                                     # most common first
    for r in d.sort_values("o").itertuples():
        noest = (not r.ref) and np.isnan(r.rr)
        t = KEY.get(fac, "") + str(r.level)
        if r.ref:
            aor, aor_b, aor_w = ("ref" if fac in KEY else "not in model"), "", ""
        elif fac not in KEY:
            aor, aor_b, aor_w = "not in model", "", ""
        elif t in mv.index:
            o = mv.loc[t]; aor, aor_b, aor_w = f"{o.OR:.2f}", ci(o.boot_lo, o.boot_hi, 2), ci(o.lo, o.hi, 2)
        else:
            aor, aor_b, aor_w = "–", "", ""
        rows.append({
            "Factor": fac, "Level": str(r.level).replace("_", " "), "k/n": f"{r.k}/{r.n:,}",
            "SAR %": f"{100 * r.sar:.1f}",
            "SAR 95% CI (bootstrap)": ci(100 * r.boot_lo, 100 * r.boot_hi), "SAR 95% CI (Wilson)": ci(100 * r.lo, 100 * r.hi),
            "Crude RR": "ref" if r.ref else ("–" if noest else f"{r.rr:.2f}"),
            "RR 95% CI (bootstrap)": "" if (r.ref or noest) else ci(r.rr_boot_lo, r.rr_boot_hi, 2),
            "RR 95% CI (Wald)": "" if (r.ref or noest) else ci(r.rr_lo, r.rr_hi, 2),
            "Adjusted OR": aor, "aOR 95% CI (bootstrap)": aor_b, "aOR 95% CI (Wald)": aor_w,
            "_k": r.k,
        })
tab = pd.DataFrame(rows)
tab.drop(columns="_k").to_csv(os.path.join(T, "si5_risk_factors.csv"), index=False)
print(tab.drop(columns="_k").to_string(index=False))

foot = (f"Denominator: {N_ELIG:,} source–contact pairs (contacts of confirmed/probable source cases whose first follow-up was on or before "
        f"{fmt_day(FU_WINDOW_END)}). Secondary case: confirmed/probable in the linelist or CT 'case confirmed'/'case validated'. Reference level: "
        "'unknown' for relationship; otherwise the most common level. Adjusted OR: single logistic regression including exposure type, "
        "setting, relationship, source province, source outcome and month listed; levels with no secondary cases are dropped ('–'). "
        "95% CI (bootstrap): 2,000 resamples of source cases, accounting for clustering of contacts within source cases (used in the "
        "text); Wilson / Wald CIs assume independent contacts. Bootstrap CIs are unreliable for levels with fewer than 5 secondary "
        "cases (shown in grey) and are 0.0–0.0 where there were none.")

# ------------------------------------------------------------------ Word
INK = RGBColor(0x01, 0x45, 0x4F); GREY = RGBColor(0x8A, 0x9A, 0x9C)
doc = Document()
sec = doc.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE
sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
    setattr(sec, side, Cm(1.5))
st = doc.styles["Normal"]; st.font.name = "Arial"; st.font.size = Pt(8)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
h = doc.add_paragraph(); r_ = h.add_run("SI-5. Risk factors for becoming a secondary case"); r_.bold = True; r_.font.size = Pt(10); r_.font.color.rgb = INK


def shade(cell, hexcol):
    tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), hexcol); tcPr.append(sh)


cols = [c for c in tab.columns if c != "_k"]
widths = [2.6, 3.6, 1.8, 1.3, 2.3, 2.3, 1.5, 2.3, 2.3, 1.9, 2.3, 2.3]
t = doc.add_table(rows=1, cols=len(cols)); t.style = "Table Grid"; t.autofit = False
lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); t._tbl.tblPr.append(lay)
for i, c in enumerate(cols):
    cell = t.rows[0].cells[i]; cell.text = ""
    rr = cell.paragraphs[0].add_run(c); rr.bold = True; rr.font.size = Pt(7.5); rr.font.color.rgb = RGBColor(255, 255, 255)
    shade(cell, "01454f")
prev = None
for ri, row in tab.iterrows():
    cells = t.add_row().cells
    sparse = row["_k"] < 5
    for i, c in enumerate(cols):
        v = row[c]
        if c == "Factor" and v == prev:
            v = ""
        cells[i].text = ""
        rr = cells[i].paragraphs[0].add_run(str(v)); rr.font.size = Pt(7.5)
        if sparse and "bootstrap" in c:
            rr.font.color.rgb = GREY
        if c == "Factor":
            rr.bold = True
        if i > 1:
            cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    prev = row["Factor"]
    if ORDER.index(row["Factor"]) % 2 == 1:
        for c_ in cells: shade(c_, "f2f7f8")
for i, w in enumerate(widths):
    t.columns[i].width = Cm(w)
for row in t.rows:
    for i, w in enumerate(widths):
        row.cells[i].width = Cm(w)
p = doc.add_paragraph(); rr = p.add_run(foot); rr.italic = True; rr.font.size = Pt(7)
path = os.path.join(OUT_ROOT, "si5_risk_factors.docx")
doc.save(path)
print("saved", path)
