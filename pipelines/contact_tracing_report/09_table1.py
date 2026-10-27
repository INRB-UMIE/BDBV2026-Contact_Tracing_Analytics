"""
Step 9: manuscript Table 1 -- key contact-tracing and epidemiological parameters.
Reads derived/ + tables/ written by steps 01, 02 and 05 (R fits must have been run for
the serial-interval / incubation rows). Writes:
  tables/table1.csv
  table1.docx                 (formatted Word table, in the output folder)

Rows (all on the main-manuscript definitions):
  Contact tracing coverage   % of confirmed/probable cases with >=1 contact listed (source cases)
  Follow-up coverage         % of contacts with a closed 21-day window (first follow-up <= cut - 21 d)
                             with a result recorded to day 21 of follow-up (cases: to detection)
  Secondary attack rate      main definition (linelist + CT 'case confirmed' + CT 'case validated')
  Relative risk by relationship, vs 'unknown' relationship
  Incubation period, serial interval   gamma fits, main analyses (05_fit_distributions_epidist.R)
  Time from source-case notification to first contact follow-up
CIs: Wilson for case-level proportions. For follow-up coverage, SAR and risk ratios the table shows the
95% bootstrap CI resampling source cases (contacts of the same source are not independent; computed in
02_analysis.py), with the Wilson / log-normal Wald CI in the notes.
"""
import os, json
import numpy as np
import pandas as pd
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from ctr_paths import D, T, OUT_ROOT, REPORT_DATE

MAX_DELAY = 60
N_BOOT = 2000
rng = np.random.default_rng(20261009)

R = json.load(open(os.path.join(D, "results.json")))
cases = pd.read_csv(os.path.join(D, "cases_annotated.csv"), low_memory=False)
pr = pd.read_csv(os.path.join(D, "pairs_annotated.csv"), low_memory=False,
                 parse_dates=["first_fu", "registration", "src_notification"])
rf = pd.read_csv(os.path.join(T, "sar_risk_factors.csv"))
gamma_path = os.path.join(T, "distribution_fits_main_gamma.csv")
gam = pd.read_csv(gamma_path) if os.path.exists(gamma_path) else None


def wilson(k, n, z=1.96):
    p = k / n; d = 1 + z * z / n; c = (p + z * z / (2 * n)) / d
    h = z * np.sqrt(p * (1 - p) / n + z * z / (4 * n * n)) / d
    return 100 * p, 100 * max(0, c - h), 100 * min(1, c + h)


def boot_by_source(df, col):
    a = df.groupby("source_case_person_id")[col].agg(["sum", "size"]).values
    i = rng.integers(0, len(a), (N_BOOT, len(a)))
    b = a[i, 0].sum(1) / a[i, 1].sum(1)
    return tuple(100 * np.percentile(b, [2.5, 97.5]))


rows = []


def add(section, measure, value, n_text="", note=""):
    rows.append(dict(Section=section, Measure=measure, Estimate=value, n=n_text, Notes=note))


# ---- contact tracing coverage
k, n = int(cases.is_source.astype(bool).sum()), len(cases)
p, lo, hi = wilson(k, n)
add("Contact tracing coverage", "Cases with >=1 contact listed", f"{p:.1f}% ({lo:.1f}-{hi:.1f})", f"{k:,}/{n:,} cases",
    "Confirmed/probable cases appearing as a source case in the CT data")
k2 = int(cases.prev_contact.astype(bool).sum())
p, lo, hi = wilson(k2, n)
add("Contact tracing coverage", "Cases already a listed contact before notification", f"{p:.1f}% ({lo:.1f}-{hi:.1f})", f"{k2:,}/{n:,} cases")

# ---- follow-up coverage
pf = pr[pr.fu_ff_eligible.astype(bool)].copy()
pf["fu_ff_complete"] = pf.fu_ff_complete.astype(bool)
k, n = int(pf.fu_ff_complete.sum()), len(pf)
p, lo, hi = wilson(k, n); blo, bhi = (100 * x for x in R["fu_complete_boot"])
add("Follow-up coverage", "Contacts fully followed up (to day 21 of follow-up)",
    f"{p:.1f}% ({blo:.1f}-{bhi:.1f})", f"{k:,}/{n:,} contacts",
    f"Closed windows: first follow-up on/before data cut - 21 d. CI: bootstrap by source case (Wilson {lo:.1f}-{hi:.1f})")

# ---- secondary attack rate
E = pr[pr.eligible.astype(bool)].copy()
E["sc"] = E.secondary_case.astype(bool)
k, n = int(E.sc.sum()), len(E)
p, lo, hi = wilson(k, n); blo, bhi = (100 * x for x in R["sar_boot"])
add("Secondary attack rate", "Overall SAR", f"{p:.1f}% ({blo:.1f}-{bhi:.1f})", f"{k:,}/{n:,} contacts",
    f"Linelist confirmed/probable or CT 'case confirmed'/'case validated'. CI: bootstrap by source case (Wilson {lo:.1f}-{hi:.1f})")

# ---- relative risk by relationship (vs unknown)
rel = rf[rf.factor == "Relationship"].copy()
order = ["unknown", "spouse", "parent", "child", "sibling", "neighbour"]
rel["o"] = rel.level.map({l: i for i, l in enumerate(order)}).fillna(99)
for r in rel.sort_values("o").itertuples():
    lab = r.level.capitalize()
    if r.ref:
        val = "1 (reference)"
    elif np.isnan(r.rr):
        val = "-"
    else:
        val = f"{r.rr:.2f} ({r.rr_boot_lo:.2f}-{r.rr_boot_hi:.2f})"
    note = f"SAR {100 * r.sar:.1f}% ({int(r.k)}/{int(r.n):,})" + ("" if r.ref or np.isnan(r.rr) else f"; Wald CI {r.rr_lo:.2f}-{r.rr_hi:.2f}")
    add("Risk ratio by relationship (vs unknown)", lab, val, note)

# ---- incubation & serial interval (gamma, main analyses)
for q in ["Incubation period", "Serial interval"]:
    if gam is None or not (gam.Quantity == q).any():
        add(q, "Gamma fit", "not available - run 05_fit_distributions_epidist.R", "")
        continue
    g = gam[gam.Quantity == q].iloc[0]
    sec = f"{q} (gamma fit, days)"
    add(sec, "Mean (95% CrI)", g["Mean (95% CrI)"], f"{int(g.n)} pairs")
    add(sec, "Median (95% CrI)", g["Median (95% CrI)"])
    add(sec, "SD (95% CrI)", g["SD (95% CrI)"])
    add(sec, "95th percentile (95% CrI)", g["95th percentile (95% CrI)"])
    add(sec, "Parameters", str(g["Gamma parameters (median, 95% CrI)"]).replace("  [", " ["),
        note="" if bool(g.converged) else "WARNING: fit flagged as not converged")

# ---- time from source-case notification to first contact follow-up
src = (pr[pr.src_is_case.astype("boolean").fillna(False)]
       .groupby("source_case_person_id").agg(first_fu=("first_fu", "min"), notif=("src_notification", "first")))
dly = (src.first_fu - src.notif).dt.days
v = dly[dly.between(-MAX_DELAY, MAX_DELAY)]
add("Time from source-case notification to first contact follow-up", "Mean (days)", f"{v.mean():.1f}",
    f"{len(v):,} source cases",
    f"Median {v.median():g} (IQR {v.quantile(.25):g}-{v.quantile(.75):g}); {int((v < 0).sum())} negative (follow-up dated before notification); "
    f"{int(dly.notna().sum() - len(v))} outside +/-{MAX_DELAY} d excluded")

tab = pd.DataFrame(rows)
tab.to_csv(os.path.join(T, "table1.csv"), index=False)
print(tab.to_string(index=False))

# ------------------------------------------------------------------ Word version
INK = RGBColor(0x01, 0x45, 0x4F)
doc = Document()
st = doc.styles["Normal"]; st.font.name = "Arial"; st.font.size = Pt(9)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
h = doc.add_paragraph(); r_ = h.add_run("Table 1. Contact tracing performance and epidemiological parameters, "
                                         f"Bundibugyo virus disease, DRC (data cut {pd.Timestamp(REPORT_DATE):%d %B %Y})")
r_.bold = True; r_.font.color.rgb = INK


def shade(cell, hexcol):
    tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), hexcol); tcPr.append(sh)


cols = ["Measure", "Estimate (95% CI/CrI)", "n", "Notes"]
widths = [4.6, 5.0, 3.2, 4.6]
t = doc.add_table(rows=1, cols=4); t.style = "Table Grid"; t.autofit = False
for i, c in enumerate(cols):
    cell = t.rows[0].cells[i]; cell.text = ""
    rr = cell.paragraphs[0].add_run(c); rr.bold = True; rr.font.color.rgb = RGBColor(255, 255, 255); rr.font.size = Pt(8.5)
    shade(cell, "01454f")
for sec, g in tab.groupby("Section", sort=False):
    cells = t.add_row().cells
    m = cells[0].merge(cells[3]); m.text = ""
    rr = m.paragraphs[0].add_run(sec); rr.bold = True; rr.font.size = Pt(8.5); rr.font.color.rgb = INK
    shade(m, "e3f0f1")
    for r in g.itertuples():
        cells = t.add_row().cells
        for i, v in enumerate([r.Measure, r.Estimate, r.n, r.Notes]):
            cells[i].text = ""
            rr = cells[i].paragraphs[0].add_run("" if pd.isna(v) else str(v)); rr.font.size = Pt(8 if i == 3 else 8.5)
            if i == 1:
                cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); t._tbl.tblPr.append(lay)
for i, w in enumerate(widths):
    t.columns[i].width = Cm(w)
for row in t.rows:
    for i, w in enumerate(widths):
        row.cells[i].width = Cm(w)
doc.add_paragraph().add_run(
    "CI: Wilson for case-level proportions; for follow-up coverage, SAR and risk ratios, 95% bootstrap CI resampling source "
    "cases (2,000 replicates; Wilson / Wald CI in notes). Incubation period and serial interval: gamma distributions, posterior median and 95% "
    "credible interval, fitted with epidist accounting for daily interval censoring and right truncation (main analysis "
    "inclusion criteria; see Methods).").font.size = Pt(7.5)
out = os.path.join(OUT_ROOT, "table1.docx")
doc.save(out)
print("saved", out)
