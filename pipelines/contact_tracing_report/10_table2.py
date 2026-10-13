"""
Step 10: manuscript Table 2 -- contact-tracing indicators by province.
Reads tables/province_table.csv and derived/sources.csv (02_analysis.py). Writes:
  tables/table2_province.csv
  table2.docx                 (formatted Word table, in the output folder)

Columns use the main-manuscript definitions:
  Cases / Deaths                     confirmed + probable linelist cases (province of the case)
  % cases with contacts listed       cases appearing as a source case in the CT data
  % cases known contact before notif.  case was a listed contact on/before own notification
  Source cases* / Contacts*          CT source persons / distinct contacts (province of the source case)
  Median contacts / source
  % contacts fully followed up       closed windows (first follow-up <= data cut - 21 d); result to day 21
  SAR % (k/n)                        main definition, contacts of confirmed/probable sources with closed windows
The province with the highest SAR (among provinces with >= 100 eligible contacts) is shown in bold.
"""
import os
import numpy as np
import pandas as pd
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from ctr_paths import D, T, OUT_ROOT, REPORT_DATE

ORDER = ["Ituri", "Nord-Kivu", "Haut-Uele", "Tshopo", "Bas-Uele", "Sud-Kivu", "National"]
MIN_N_BOLD = 100

pt = pd.read_csv(os.path.join(T, "province_table.csv")).set_index("area")
src = pd.read_csv(os.path.join(D, "sources.csv"))
order = [a for a in ORDER if a in pt.index] + [a for a in pt.index if a not in ORDER]
pt = pt.reindex(order)

fmt_int = lambda x: "-" if pd.isna(x) else f"{int(x):,}"
fmt1 = lambda x: "-" if pd.isna(x) else f"{x:.1f}"
fmt0 = lambda x: "-" if pd.isna(x) else f"{x:.0f}"
fmtg = lambda x: "-" if pd.isna(x) else f"{x:g}"
out = pd.DataFrame({
    "Province": pt.index,
    "Cases": pt.cases.map(fmt_int).values,
    "Deaths": pt.deaths.map(fmt_int).values,
    "% cases with contacts listed": pt.pct_cases_source.map(fmt1).values,
    "% cases known contact before notif.": pt.pct_cases_prev_contact.map(fmt1).values,
    "Source cases*": pt.source_cases.map(fmt_int).values,
    "Contacts*": pt.contacts.map(fmt_int).values,
    "Median contacts / source": pt.med_contacts.map(fmtg).values,
    "% contacts fully followed up": pt.pct_fu_complete.map(fmt0).values,
    "SAR % (k/n)": [f"{100 * k / n:.1f} ({int(k)}/{int(n):,})" if n else "-" for k, n in zip(pt.sar_k, pt.sar_n)],
})
out.to_csv(os.path.join(T, "table2_province.csv"), index=False)
print(out.to_string(index=False))

prov = pt.drop(index="National", errors="ignore")
elig = prov[prov.sar_n >= MIN_N_BOLD]
bold_row = (elig.sar_k / elig.sar_n).idxmax() if len(elig) else None

n_kasai = int((src.province == "Kasai").sum())
n_unloc = int(src.province.isna().sum())
n_other = {p: int((src.province == p).sum()) for p in sorted(set(src.province.dropna()) - set(ORDER) - {"Kasai"})}
other_txt = "".join(f", {n} in {p}" for p, n in n_other.items())
foot = (f"*Province of the source case. National totals include {n_kasai} source persons with Kasai (Bulape) alert-ID prefixes, "
        f"{n_unloc} with unlocatable IDs{other_txt}. Follow-up: contacts whose first follow-up was on or before "
        f"{pd.Timestamp(REPORT_DATE) - pd.Timedelta(days=21):%d %b %Y} (21 days before the data cut), with a result recorded to day 21 "
        "of follow-up (secondary cases: to detection). SAR: contacts of confirmed/probable source cases in that group; secondary case = "
        "confirmed/probable in the linelist or CT 'case confirmed'/'case validated'."
        + (f" Bold: highest SAR among provinces with ≥{MIN_N_BOLD} eligible contacts." if bold_row else ""))

# ------------------------------------------------------------------ Word
INK = RGBColor(0x01, 0x45, 0x4F)
doc = Document()
sec = doc.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE
sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
    setattr(sec, side, Cm(1.8))
st = doc.styles["Normal"]; st.font.name = "Arial"; st.font.size = Pt(9)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
h = doc.add_paragraph()
r_ = h.add_run(f"Table 2. Contact-tracing indicators by province (data cut {pd.Timestamp(REPORT_DATE):%d %B %Y})")
r_.bold = True; r_.font.color.rgb = INK


def shade(cell, hexcol):
    tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), hexcol); tcPr.append(sh)


widths = [2.6, 2.0, 2.0, 2.6, 2.7, 2.2, 2.2, 2.4, 2.6, 3.4]
t = doc.add_table(rows=1, cols=len(out.columns)); t.style = "Table Grid"; t.autofit = False
lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); t._tbl.tblPr.append(lay)
for i, c in enumerate(out.columns):
    cell = t.rows[0].cells[i]; cell.text = ""
    rr = cell.paragraphs[0].add_run(c); rr.bold = True; rr.font.size = Pt(8.5); rr.font.color.rgb = RGBColor(255, 255, 255)
    shade(cell, "01454f")
for ri, row in enumerate(out.itertuples(index=False)):
    cells = t.add_row().cells
    is_bold = row[0] == bold_row or row[0] == "National"
    for i, v in enumerate(row):
        cells[i].text = ""
        rr = cells[i].paragraphs[0].add_run(str(v)); rr.font.size = Pt(8.5); rr.bold = is_bold
        if i > 0:
            cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if row[0] == "National":
        for c in cells: shade(c, "e3f0f1")
    elif ri % 2 == 1:
        for c in cells: shade(c, "f2f7f8")
for i, w in enumerate(widths):
    t.columns[i].width = Cm(w)
for row in t.rows:
    for i, w in enumerate(widths):
        row.cells[i].width = Cm(w)
p = doc.add_paragraph(); rr = p.add_run(foot); rr.italic = True; rr.font.size = Pt(7.5)
path = os.path.join(OUT_ROOT, "table2.docx")
doc.save(path)
print("saved", path)
